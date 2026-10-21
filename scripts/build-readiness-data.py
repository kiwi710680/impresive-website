# Usage: python -X utf8 scripts/build-readiness-data.py "<path to mSCDM .docx>" assets/js/readiness-data.js
import docx, json, re, sys
src, out = sys.argv[1], sys.argv[2]
d = docx.Document(src)

def cells(row):
    seen, vals = set(), []
    for c in row.cells:
        if id(c._tc) in seen: continue
        seen.add(id(c._tc)); vals.append(c.text.strip())
    return vals

def clean_name(n):
    return re.sub(r'\s*\d+$', '', n.strip())

def clean_type(t):
    return re.sub(r'\)\d+$', ')', t.strip())

def tidy(t):
    t = re.sub(r'[ \t]+', ' ', t)
    return re.sub(r'\s*\n\s*', '\n', t).strip()

# table order in the document after the table of contents
order = [
 ("enrollment", "Enrollment", "core"), ("demographic", "Demographic", "core"), ("encounter", "Encounter", "core"),
 ("diagnosis", "Diagnosis", "core"), ("procedure", "Procedure", "core"), ("death", "Death", "core"),
 ("cause-of-death", "Cause of Death", "core"), ("prescribing", "Prescribing", "core"),
 ("facility", "Facility", "supporting"), ("provider", "Provider", "supporting"), ("drug-strength", "Drug_Strength", "supporting")]

# Essential = required for the table to be usable (identifiers, linking keys,
# the core date and code fields). Everything else is supplementary.
essential = {
 "enrollment": {"PatID", "Enr_Start", "Enr_End"},
 "demographic": {"PatID", "Birth_Date", "Sex"},
 "encounter": {"PatID", "EncounterID", "ADate", "DDate", "EncType"},
 "diagnosis": {"PatID", "EncounterID", "ADate", "DX", "Dx_Codetype"},
 "procedure": {"PatID", "EncounterID", "ADate", "PX", "PX_CodeType"},
 "death": {"PatID", "DeathDt"},
 "cause-of-death": {"PatID", "COD", "CodeType"},
 "prescribing": {"PatID", "OrderDate", "Rx_CodeType", "Rx", "RxSup"},
 "facility": {"FacilityID", "Facility_year"},
 "provider": {"ProviderID", "Specialty"},
 "drug-strength": {"Rx", "Rx_CodeType", "Ingred_ATC"},
}

toc = {}
for r in d.tables[0].rows[1:]:
    v = cells(r)
    toc[re.sub(r'^\d+\.\s*', '', v[0]).strip()] = tidy(v[2])

tables = []
for (key, label, tier), t in zip(order, d.tables[1:]):
    rows = [cells(r) for r in t.rows]
    head = rows[0]
    variables = []
    for v in rows[1:]:
        name = clean_name(v[0])
        variables.append({
            "name": name,
            "type": clean_type(v[1]),
            "values": tidy(v[2]),
            "missingAllowed": tidy(v[3]),
            "definition": tidy(v[4]),
            "example": tidy(v[5]) if len(v) > 5 else "",
            "essential": name in essential[key],
        })
    missing = essential[key] - {x["name"] for x in variables}
    assert not missing, (key, missing)
    desc = toc.get(label) or toc.get(label.replace("_", " "))
    desc = re.sub(r'(?<=[A-Za-z.])1(?=[\s.]|$)', '', desc or "")   # drop footnote markers
    tables.append({"id": key, "label": label, "tier": tier, "description": desc, "variables": variables})

data = {"source": "Modified Sentinel Common Data Model (mSCDM) 8.1.0 for PHDcenter", "tables": tables}
js = "/* Generated from \"Sentinel Common Data Model v8.1.0 modify for PHDcenter.docx\". */\nwindow.IMPRESIVE_READINESS = " + json.dumps(data, ensure_ascii=False, indent=1) + ";\n"
open(out, "w", encoding="utf8").write(js)
print(len(tables), sum(len(t["variables"]) for t in tables), sum(v["essential"] for t in tables for v in t["variables"]))
for t in tables: print(t["id"], [v["name"] + ("*" if v["essential"] else "") for v in t["variables"]], "|", t["description"][:90])
